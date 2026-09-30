# -*- coding: utf-8 -*-
"""Parse DOCX properties into a stable, reviewable document model.

The paragraph list retains top-level body indices for existing reports. Table
paragraphs and section stories have separate locations; no page number or visual
position is inferred from XML alone.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.table import Table

CJK_RE = re.compile(r"[一-鿿]")
ASCII_RE = re.compile(r"[A-Za-z0-9]")

CN_FONT_SIZE = {
    "初号": 42, "小初": 36, "一号": 26, "小一": 24, "二号": 22, "小二": 18,
    "三号": 16, "小三": 15, "四号": 14, "小四": 12, "五号": 10.5, "小五": 9,
}


def normalize_font(name: str | None) -> str:
    if not name:
        return ""
    return re.sub(r"[\s_\-]+", "", str(name)).lower()


def _align_name(align) -> str | None:
    if align is None:
        return None
    mapping = {
        WD_ALIGN_PARAGRAPH.LEFT: "LEFT",
        WD_ALIGN_PARAGRAPH.CENTER: "CENTER",
        WD_ALIGN_PARAGRAPH.RIGHT: "RIGHT",
        WD_ALIGN_PARAGRAPH.JUSTIFY: "JUSTIFY",
    }
    return mapping.get(align, str(align))


@dataclass
class RunInfo:
    text: str
    font_ea: str | None
    font_ascii: str | None
    size_pt: float | None
    bold: bool | None
    char_spacing_pt: float | None
    sources: dict[str, str] = field(default_factory=dict)
    unresolved: dict[str, str] = field(default_factory=dict)


@dataclass
class ParaInfo:
    index: int
    text: str
    style_name: str
    alignment: str | None
    line_pt: float | None
    line_multiple: float | None
    line_rule: str | None
    first_line_chars: int | None
    left_chars: int | None
    right_chars: int | None
    runs: list[RunInfo] = field(default_factory=list)
    location: str = ""
    sources: dict[str, str] = field(default_factory=dict)

    @property
    def is_blank(self) -> bool:
        return self.text.strip() == ""

    def content_runs(self) -> list[RunInfo]:
        return [r for r in self.runs if r.text.strip()]

    def main_run(self) -> RunInfo | None:
        crs = self.content_runs()
        return max(crs, key=lambda r: len(r.text)) if crs else None

    def main_size(self) -> float | None:
        r = self.main_run()
        return r.size_pt if r else None

    def snippet(self, n: int = 14) -> str:
        t = self.text.strip().replace("\n", " ")
        return t[:n] + ("…" if len(t) > n else "")


@dataclass
class FooterPara:
    text: str
    alignment: str | None
    has_page_field: bool
    runs: list[RunInfo] = field(default_factory=list)
    location: str = "页脚"
    section_index: int = 0
    kind: str = "default"
    left_chars: int | None = None
    right_chars: int | None = None


@dataclass
class SectionInfo:
    top_cm: float | None
    bottom_cm: float | None
    left_cm: float | None
    right_cm: float | None
    footer_distance_cm: float | None
    different_first_page_header_footer: bool = False


@dataclass
class DocModel:
    path: str
    paragraphs: list[ParaInfo]
    sections: list[SectionInfo]
    footer_paras: list[FooterPara]
    table_paragraphs: list[ParaInfo] = field(default_factory=list)
    header_paras: list[FooterPara] = field(default_factory=list)
    body_order: list[str] = field(default_factory=list)
    unchecked_parts: list[str] = field(default_factory=list)
    format_only_unchecked_parts: list[str] = field(default_factory=list)
    story_links: dict[str, str] = field(default_factory=dict)
    odd_even_header_footer: bool = False


def _style_chain(style):
    seen = set()
    while style is not None and id(style.element) not in seen:
        seen.add(id(style.element))
        yield style
        style = style.base_style


def _run_values(rpr) -> dict:
    if rpr is None:
        return {}
    values: dict = {}
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is not None:
        for key, attrs in {
            "font_ea": ("w:eastAsia",),
            "font_ascii": ("w:ascii", "w:hAnsi"),
        }.items():
            value = next((fonts.get(qn(attr)) for attr in attrs if fonts.get(qn(attr))), None)
            if value is not None:
                values[key] = value
        theme_attrs = {
            "font_ea_theme": ("w:eastAsiaTheme",),
            "font_ascii_theme": ("w:asciiTheme", "w:hAnsiTheme"),
        }
        for key, attrs in theme_attrs.items():
            value = next((fonts.get(qn(attr)) for attr in attrs if fonts.get(qn(attr))), None)
            if value is not None:
                values[key] = value
    size = rpr.find(qn("w:sz"))
    if size is not None and size.get(qn("w:val")):
        values["size_pt"] = int(size.get(qn("w:val"))) / 2
    bold = rpr.find(qn("w:b"))
    if bold is not None:
        values["bold"] = bold.get(qn("w:val"), "1").lower() not in {"0", "false", "off", "no"}
    spacing = rpr.find(qn("w:spacing"))
    if spacing is not None and spacing.get(qn("w:val")) is not None:
        values["char_spacing_pt"] = int(spacing.get(qn("w:val"))) / 20
    return values


def _paragraph_values(ppr) -> dict:
    if ppr is None:
        return {}
    values: dict = {}
    spacing = ppr.find(qn("w:spacing"))
    if spacing is not None and spacing.get(qn("w:line")):
        line_rule = spacing.get(qn("w:lineRule")) or "auto"
        raw_line = int(spacing.get(qn("w:line")))
        values["line_rule"] = line_rule
        if line_rule == "auto":
            values["line_multiple"] = raw_line / 240
        else:
            values["line_pt"] = raw_line / 20
    indent = ppr.find(qn("w:ind"))
    if indent is not None:
        for attr, name in (("w:firstLineChars", "first_line_chars"),
                           ("w:leftChars", "left_chars"),
                           ("w:rightChars", "right_chars")):
            if indent.get(qn(attr)) is not None:
                values[name] = int(indent.get(qn(attr)))
    jc = ppr.find(qn("w:jc"))
    if jc is not None:
        mapping = {"left": "LEFT", "center": "CENTER", "right": "RIGHT",
                   "both": "JUSTIFY", "start": "LEFT", "end": "RIGHT"}
        values["alignment"] = mapping.get(jc.get(qn("w:val")), None)
    return values


def _defaults(styles, *, paragraph: bool):
    if styles is None:
        return None
    defaults = styles.element.find(qn("w:docDefaults"))
    if defaults is None:
        return None
    group = "w:pPrDefault" if paragraph else "w:rPrDefault"
    prop = "w:pPr" if paragraph else "w:rPr"
    wrapper = defaults.find(qn(group))
    return wrapper.find(qn(prop)) if wrapper is not None else None


def _effective(layers: list[tuple[str, dict]], keys: tuple[str, ...]) -> tuple[dict, dict]:
    values = {key: None for key in keys}
    sources: dict[str, str] = {}
    for source, layer in layers:
        for key in keys:
            if values[key] is None and key in layer and layer[key] is not None:
                values[key] = layer[key]
                sources[key] = source
    return values, sources


def _parse_run(run, paragraph_style=None, styles=None) -> RunInfo:
    layers = [("run", _run_values(run._element.rPr))]
    if run.style is not None:
        for style in _style_chain(run.style):
            layers.append((f"character_style:{style.name}", _run_values(style.element.rPr)))
    if paragraph_style is not None:
        for style in _style_chain(paragraph_style):
            layers.append((f"paragraph_style:{style.name}", _run_values(style.element.rPr)))
    layers.append(("doc_defaults", _run_values(_defaults(styles, paragraph=False))))
    keys = ("font_ea", "font_ascii", "size_pt", "bold", "char_spacing_pt")
    values, sources = _effective(layers, keys)
    unresolved: dict[str, str] = {}
    for font_key, theme_key in (("font_ea", "font_ea_theme"),
                                ("font_ascii", "font_ascii_theme")):
        for source, layer in layers:
            # A theme reference at the controlling layer needs theme resolution.
            # A direct font at a more specific layer takes precedence over a
            # theme reference inherited from a less specific style.
            if font_key in layer:
                break
            if theme_key in layer:
                values[font_key] = None
                sources.pop(font_key, None)
                unresolved[font_key] = f"{source}:{layer[theme_key]}"
                break
    if values["bold"] is None:
        values["bold"] = False
        sources["bold"] = "implicit_default"
    return RunInfo(text=run.text or "", sources=sources,
                   unresolved=unresolved, **values)


def _parse_paragraph(p, index: int, styles=None, location: str = "") -> ParaInfo:
    layers = [("paragraph", _paragraph_values(p._p.pPr))]
    if p.style is not None:
        for style in _style_chain(p.style):
            layers.append((f"paragraph_style:{style.name}", _paragraph_values(style.element.pPr)))
    layers.append(("doc_defaults", _paragraph_values(_defaults(styles, paragraph=True))))
    keys = ("alignment", "line_pt", "line_multiple", "line_rule", "first_line_chars", "left_chars", "right_chars")
    values, sources = _effective(layers, keys)
    # w:line uses 1/240 of a line for auto, and 1/20 pt for exact/atLeast.
    # A direct multiple value must not inherit a point value from a style.
    if values["line_rule"] == "auto":
        values["line_pt"] = None
        sources.pop("line_pt", None)
    elif values["line_rule"] in {"exact", "atLeast"}:
        values["line_multiple"] = None
        sources.pop("line_multiple", None)
    return ParaInfo(
        index=index, text=p.text or "",
        style_name=p.style.name if p.style is not None else "",
        runs=[_parse_run(r, p.style, styles) for r in p.runs],
        location=location or f"body/p{index + 1}", sources=sources, **values,
    )


def _parse_story_para(p, location: str, section_index: int, kind: str, styles) -> FooterPara:
    el = p._p
    full_text = "".join(t.text or "" for t in el.findall(".//" + qn("w:t")))
    has_page = any(re.search(r"\bPAGE\b", it.text or "", re.IGNORECASE) for it in el.findall(".//" + qn("w:instrText")))
    has_page = has_page or any(
        re.search(r"\bPAGE\b", f.get(qn("w:instr")) or "", re.IGNORECASE) for f in el.findall(".//" + qn("w:fldSimple")))
    parsed = _parse_paragraph(p, -1, styles, location)
    return FooterPara(
        text=full_text, alignment=parsed.alignment, has_page_field=has_page,
        runs=parsed.runs, location=location, section_index=section_index, kind=kind,
        left_chars=parsed.left_chars, right_chars=parsed.right_chars,
    )


def parse_document(path: str) -> DocModel:
    doc = Document(path)
    styles = doc.styles
    paragraphs = [
        _parse_paragraph(p, i, styles, f"body/p{i + 1}")
        for i, p in enumerate(doc.paragraphs)
    ]
    body_order = []
    table_paragraphs = []
    table_style_notes = []
    nested_table_notes = []
    body_index = table_index = 0
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            body_index += 1
            body_order.append(f"body/p{body_index}")
        elif child.tag == qn("w:tbl"):
            table_index += 1
            table = Table(child, doc._body)
            body_order.append(f"table{table_index}")
            tbl_pr = child.find(qn("w:tblPr"))
            tbl_style = tbl_pr.find(qn("w:tblStyle")) if tbl_pr is not None else None
            if tbl_style is not None:
                style_id = tbl_style.get(qn("w:val")) or "未命名"
                table_style_notes.append(
                    f"table{table_index} 使用表格样式 {style_id}，样式属性未解析及检查")
            for row_index, row in enumerate(table.rows, 1):
                for cell_index, cell in enumerate(row.cells, 1):
                    cell_loc = f"table{table_index}/r{row_index}c{cell_index}"
                    if cell.tables:
                        nested_table_notes.append(
                            f"{cell_loc} 含 {len(cell.tables)} 个嵌套表格，内部文字未解析及检查")
                    for para_index, p in enumerate(cell.paragraphs, 1):
                        loc = f"table{table_index}/r{row_index}c{cell_index}/p{para_index}"
                        table_paragraphs.append(_parse_paragraph(p, -1, styles, loc))
    sections: list[SectionInfo] = []
    footer_paras: list[FooterPara] = []
    header_paras: list[FooterPara] = []
    even_enabled = bool(doc.settings.odd_and_even_pages_header_footer)
    story_links: dict[str, str] = {}
    labels = {"default": "普通", "first": "首页", "even": "偶数页"}
    for section_index, section in enumerate(doc.sections, 1):
        def cm(value):
            return round(value.cm, 3) if value is not None else None
        sections.append(SectionInfo(
            top_cm=cm(section.top_margin), bottom_cm=cm(section.bottom_margin),
            left_cm=cm(section.left_margin), right_cm=cm(section.right_margin),
            footer_distance_cm=cm(section.footer_distance),
            different_first_page_header_footer=bool(section.different_first_page_header_footer),
        ))
        for kind in ("default", "first", "even"):
            if kind == "first" and not section.different_first_page_header_footer:
                continue
            if kind == "even" and not even_enabled:
                continue
            for part, target in (("footer", footer_paras), ("header", header_paras)):
                attr = part if kind == "default" else f"{kind}_page_{part}"
                story = getattr(section, attr)
                if story.is_linked_to_previous:
                    key = f"section{section_index}/{kind}/{part}"
                    story_links[key] = f"section{section_index - 1}/{kind}/{part}" if section_index > 1 else "未定义"
                    continue  # 已在前一节检查；记录链接供报告追溯
                label = "页脚" if part == "footer" else "页眉"
                for para_index, p in enumerate(story.paragraphs, 1):
                    loc = f"第{section_index}节{labels[kind]}{label}第{para_index}段"
                    target.append(_parse_story_para(p, loc, section_index, kind, styles))
    unchecked_parts = []
    table_content = sum(not p.is_blank for p in table_paragraphs)
    if table_content:
        unchecked_parts.append(f"表格内 {table_content} 个非空段落已解析，段落规则尚未应用")
    format_only_notes = list(table_style_notes)
    unchecked_parts.extend(table_style_notes)
    unchecked_parts.extend(nested_table_notes)
    for para in paragraphs + table_paragraphs + footer_paras + header_paras:
        for run_index, run in enumerate(para.runs, 1):
            if not run.text.strip():
                continue
            for font_key, origin in run.unresolved.items():
                if font_key == "font_ea" and not CJK_RE.search(run.text):
                    continue
                if font_key == "font_ascii" and not ASCII_RE.search(run.text):
                    continue
                note = f"{para.location}/run{run_index} 主题字体 {font_key}（{origin}）未解析及检查"
                unchecked_parts.append(note)
                format_only_notes.append(note)
    for para in paragraphs + table_paragraphs:
        if para.text and para.text != "".join(run.text for run in para.runs):
            unchecked_parts.append(f"{para.location} 含普通 run 列表之外的文字，格式规则可能未覆盖")
    header_content = sum(bool(p.text.strip()) for p in header_paras)
    if header_content:
        unchecked_parts.append(f"页眉内 {header_content} 个非空段落已解析，格式和错字规则尚未应用")
    return DocModel(
        path=path, paragraphs=paragraphs, sections=sections,
        footer_paras=footer_paras,
        table_paragraphs=table_paragraphs, header_paras=header_paras,
        body_order=body_order, unchecked_parts=unchecked_parts,
        format_only_unchecked_parts=format_only_notes, story_links=story_links,
        odd_even_header_footer=even_enabled,
    )
