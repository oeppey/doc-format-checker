"""Extract observed DOCX values as editable candidate rules, never as confirmed standards."""
from __future__ import annotations

import re

from docx import Document

from .docx_parser import parse_document
from .role_mapper import map_roles

CJK = re.compile(r"[一-鿿]")
LATIN = re.compile(r"[A-Za-z0-9]")


def extract_candidates(path: str) -> dict:
    model = parse_document(path)
    doc = Document(path)
    roles, _ = map_roles(model)
    candidates = []
    mixed = []
    def add(rid, expected, *, role=None, source="", location="", count=1):
        scope = {"part": "body", "role": role} if role else (
            {"part": "footer-all"} if rid.startswith("page_number.") else {})
        candidates.append({"id": rid, "scope": scope, "expected": expected,
                           "severity": "error", "required": False, "enabled": False,
                           "observed": {"source": source, "location": location, "count": count}})
    def unique(rid, items, make, *, role=None):
        values = [(value, source, location) for value, source, location in items if value is not None]
        if not values:
            return
        distinct = {str(value) for value, _, _ in values}
        if len(distinct) != 1:
            mixed.append({"id": rid, "role": role, "values": sorted(distinct)[:8],
                          "message": "样本文档中存在多种值，请人工选择标准"})
            return
        value, source, location = values[0]
        add(rid, make(value), role=role, source=source, location=location, count=len(values))

    unique("page.size",
           [((round(s.page_width.mm, 1), round(s.page_height.mm, 1)),
             "section", f"第{i}节")
            for i, s in enumerate(doc.sections, 1)],
           lambda size: {"width": size[0], "height": size[1], "unit": "mm"})
    for attr, rid in (("top_cm", "page.margin_top"), ("bottom_cm", "page.margin_bottom"),
                      ("left_cm", "page.margin_left"), ("right_cm", "page.margin_right"),
                      ("footer_distance_cm", "page.footer_distance")):
        unique(rid, [(getattr(s, attr), "section", f"第{i}节")
                     for i, s in enumerate(model.sections, 1)],
               lambda v: {"value": v, "unit": "cm"})
    for role in ("title", "heading1", "heading2", "heading3", "heading4",
                 "body", "wenhao", "secrecy", "attachment", "signature", "date"):
        ps = [model.paragraphs[i] for i in roles.get(role, []) if not model.paragraphs[i].is_blank]
        if not ps:
            continue
        for key, rid, present in (
            ("font_ea", "font.cjk", CJK.search), ("font_ascii", "font.latin", LATIN.search),
            ("size_pt", "font.size", lambda _: True),
            ("bold", "font.bold", lambda _: True),
            ("char_spacing_pt", "font.character_spacing", lambda _: True),
        ):
            items = []
            for p in ps:
                for r in p.content_runs():
                    if present(r.text):
                        value = getattr(r, key)
                        if key == "char_spacing_pt" and value is None:
                            value = 0.0
                        if key in r.unresolved:
                            continue
                        items.append((value, r.sources.get(key, "未显式设置"), p.location))
            if rid in {"font.cjk", "font.latin"}:
                make = lambda v: {"value": v, "aliases": []}
            elif rid == "font.bold":
                make = lambda v: {"value": v}
            else:
                make = lambda v: {"value": v, "unit": "pt"}
            unique(rid, items, make, role=role)
        unique("paragraph.alignment",
               [(p.alignment.lower(), p.sources.get("alignment", "默认值"), p.location)
                for p in ps if p.alignment in {"LEFT", "CENTER", "RIGHT"}],
               lambda v: {"value": v}, role=role)
        unique("paragraph.line_spacing",
               [((p.line_rule, p.line_multiple if p.line_rule == "auto" else p.line_pt),
                 p.sources.get("line_rule", "未设置"), p.location)
                for p in ps if p.line_rule and
                (p.line_multiple is not None or p.line_pt is not None)],
               lambda pair: ({"mode": "multiple", "value": pair[1]} if pair[0] == "auto"
                             else {"mode": "at_least" if pair[0] == "atLeast" else "exact",
                                   "value": pair[1], "unit": "pt"}), role=role)
        unique("paragraph.first_line_indent",
               [(p.first_line_chars / 100, p.sources.get("first_line_chars", "段落属性"), p.location)
                for p in ps if p.first_line_chars is not None],
               lambda v: {"value": v, "unit": "char"}, role=role)
    add("page_number.different_odd_even", {"value": model.odd_even_header_footer},
        source="settings.xml", location="文档设置")
    for rid, kinds, default_direction in (
        ("page_number.odd_padding", {"default", "first"}, "right"),
        ("page_number.even_padding", {"even"}, "left"),
    ):
        observed = []
        unsupported = []
        if model.odd_even_header_footer:
            for p in model.footer_paras:
                if p.kind not in kinds or not p.has_page_field:
                    continue
                if p.kind == "first" and p.section_index > 1:
                    continue
                left, right = p.left_chars or 0, p.right_chars or 0
                if (left and right) or any(v < 0 or v > 1000 or v % 100 for v in (left, right)):
                    unsupported.append(p.location)
                    continue
                direction = "left" if left else "right" if right else "none"
                observed.append(((direction, (left or right) // 100),
                                 "页脚段落字符缩进", p.location))
        if unsupported:
            mixed.append({"id": rid, "values": [], "locations": unsupported,
                          "message": "页码段落缩进不能唯一表示为 0–10 个整字，请人工设置标准"})
        elif observed:
            unique(rid, observed,
                   lambda pair: {"direction": pair[0], "value": pair[1], "unit": "char"})
        else:
            add(rid, {"direction": default_direction, "value": 1, "unit": "char"},
                source="建议缺省（非样本提取）", location="请用户确认")
    footer = [p for p in model.footer_paras if p.has_page_field]
    if footer:
        add("page_number.field", {"value": "PAGE"}, source="页码域", location=footer[0].location,
            count=len(footer))
        unique("page_number.alignment",
               [(p.alignment.lower(), "footer", p.location)
                for p in footer if p.alignment in {"LEFT", "CENTER", "RIGHT"}],
               lambda v: {"value": v})
        for key, rid in (("font_ea", "page_number.font_cjk"),
                         ("size_pt", "page_number.font_size")):
            items = [(getattr(r, key), r.sources.get(key, "未显式设置"), p.location)
                     for p in footer for r in p.runs if r.text.strip()]
            make = (lambda v: {"value": v, "aliases": []}) if key == "font_ea" else (
                lambda v: {"value": v, "unit": "pt"})
            unique(rid, items, make)
    return {
        "candidates": candidates,
        "mixed": mixed,
        "roles": {role: len(indices) for role, indices in roles.items()},
        "unchecked_parts": model.unchecked_parts,
        "notice": "以上仅为样本文档观测值；勾选并确认后才成为模板标准。未出现的结构不代表不需要检查。",
    }


