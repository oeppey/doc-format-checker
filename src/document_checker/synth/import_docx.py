# -*- coding: utf-8 -*-
"""导入公开公文：抽取真实内容 → 重排为目标格式（呈报稿大字版/小字版）。

抽取规则（面向两类常见公文结构）：
  - 章条式（管理办法/办法类）：首个非空段为标题；"第X章 xx" → 一级标题"X、xx"；
    "（一）"→ 二级标题；"1."→ 三级标题；其余（含"第X条"）→ 正文
  - 条款式（通知/报告类）：同上，另识别"一、"一级标题
  - 文末识别署名/成文日期；公开件不带秘级，secrecy=False

注意：表格、图片、页眉页脚中的内容暂不导入（记日志提示）。
"""
from __future__ import annotations

import re

from docx import Document

CHAPTER_RE = re.compile(r"^第\s*([一二三四五六七八九十百零]+)\s*章\s*[　 ]*(.*)$")
H1_RE = re.compile(r"^[一二三四五六七八九十百]+、")
H2_RE = re.compile(r"^（[一二三四五六七八九十百]+）")
H3_RE = re.compile(r"^\d+\.")
DATE_RE = re.compile(r"^\d{4}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日$")
# 落款单位：以 部/厅/局/委/室/办/院/中心/公司/政府 等结尾
ORG_RE = re.compile(r"(部|厅|局|委|室|办|所|院|中心|公司|政府|科协|联合会|总社)$")
NOISE_RE = re.compile(r"^(附件下载|扫一扫|（此件.*公开）|（联系人.*）)$")


def extract_content(src_path: str, default_signature: str = "XX市纪委监委XX室",
                    default_date: str = "2026年9月18日") -> dict:
    doc = Document(src_path)
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    paras = [t for t in paras if not NOISE_RE.match(t)]
    if not paras:
        raise ValueError(f"{src_path} 中没有可用文本")

    # 文末署名/日期（若有）
    signature, date = default_signature, default_date
    if paras and DATE_RE.match(paras[-1]):
        date = paras.pop(-1)
        if paras and ORG_RE.search(paras[-1]) and len(paras[-1]) <= 40:
            signature = paras.pop(-1)

    # 去掉文首的附件引导行（"附件："/"附件：xx"）
    while paras and re.match(r"^附件\s*[:：]", paras[0]):
        paras.pop(0)
    if not paras:
        raise ValueError(f"{src_path} 中没有可用文本")

    title = paras.pop(0)
    blocks: list[tuple[str, str]] = []
    skipped = 0
    for t in paras:
        m = CHAPTER_RE.match(t)
        if m:  # 第X章 xx → X、xx（章名内部空格一并去除）
            name = re.sub(r"\s+", "", m.group(2)) or "概述"
            blocks.append(("h1", f"{m.group(1)}、{name}"))
        elif H1_RE.match(t):
            blocks.append(("h1", t))
        elif H2_RE.match(t):
            blocks.append(("h2", t))
        elif H3_RE.match(t) and len(t) <= 40:
            blocks.append(("h3", t))
        elif t.startswith("附件：") or t.startswith("附件:"):
            skipped += 1  # 附件说明由重排模板统一生成
        else:
            blocks.append(("body", t))
    if len(doc.tables) or len(doc.inline_shapes):
        skipped += len(doc.tables) + len(doc.inline_shapes)

    return {
        "title": title,
        "blocks": blocks,
        "attachment": None,
        "signature": signature,
        "date": date,
        "secrecy": False,
        "_src": src_path,
        "_skipped": skipped,
    }
