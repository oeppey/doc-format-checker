"""OOXML-only checks for odd/even PAGE-field character indents."""
from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from document_checker.checkers.structure import page_number_padding
from document_checker.docx_parser import parse_document
from document_checker.standardization import compile_rules
from document_checker.template_extractor import extract_candidates


def _page_field(paragraph, *, left=0, right=0):
    ind = paragraph._p.get_or_add_pPr().find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        paragraph._p.get_or_add_pPr().append(ind)
    if left:
        ind.set(qn("w:leftChars"), str(left * 100))
    if right:
        ind.set(qn("w:rightChars"), str(right * 100))
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def _doc(tmp_path, *, first=False, linked=False, missing_even=False):
    path = tmp_path / "padding.docx"
    doc = Document()
    doc.add_paragraph("正文")
    doc.settings.odd_and_even_pages_header_footer = True
    sec = doc.sections[0]
    if first:
        sec.different_first_page_header_footer = True
        _page_field(sec.first_page_footer.paragraphs[0], right=1)
    _page_field(sec.footer.paragraphs[0], right=1)
    even_footer = sec.even_page_footer
    even_footer.is_linked_to_previous = False
    if not missing_even:
        _page_field(even_footer.paragraphs[0], left=1)
    if linked:
        doc.add_section(WD_SECTION.NEW_PAGE)
    doc.save(path)
    return path


def test_valid_character_padding_and_linked_footer(tmp_path):
    model = parse_document(str(_doc(tmp_path, first=True, linked=True)))
    assert model.sections[0].different_first_page_header_footer
    assert model.story_links["section2/even/footer"] == "section1/even/footer"
    assert page_number_padding(model, {}, {"kind": "odd", "direction": "right", "chars": 1}) == []
    assert page_number_padding(model, {}, {"kind": "even", "direction": "left", "chars": 1}) == []
    assert model.footer_paras[0].right_chars == 100
    candidates = {r["id"]: r for r in extract_candidates(model.path)["candidates"]}
    assert candidates["page_number.odd_padding"]["expected"] == {
        "direction": "right", "value": 1, "unit": "char"}
    assert candidates["page_number.even_padding"]["observed"]["source"] == "页脚段落字符缩进"


def test_wrong_side_and_missing_page_field_are_reported(tmp_path):
    model = parse_document(str(_doc(tmp_path, missing_even=True)))
    odd = page_number_padding(model, {}, {"kind": "odd", "direction": "left", "chars": 1})
    even = page_number_padding(model, {}, {"kind": "even", "direction": "left", "chars": 1})
    assert len(odd) == 1 and "普通页脚" in odd[0].location
    assert odd[0].actual == "左空 0 字、右空 1 字"
    assert len(even) == 1 and "PAGE" in even[0].expected


def test_first_page_footer_is_checked(tmp_path):
    model = parse_document(str(_doc(tmp_path, first=True)))
    first = next(p for p in model.footer_paras if p.kind == "first")
    first.right_chars = 0
    findings = page_number_padding(model, {}, {"kind": "odd", "direction": "right", "chars": 1})
    assert len(findings) == 1 and "首页页脚" in findings[0].location


def test_later_section_first_footer_parity_is_unchecked(tmp_path):
    path = _doc(tmp_path, first=True, linked=True)
    doc = Document(path)
    sec2 = doc.sections[1]
    sec2.different_first_page_header_footer = True
    first = sec2.first_page_footer
    first.is_linked_to_previous = False
    _page_field(first.paragraphs[0], left=1)
    doc.save(path)
    model = parse_document(str(path))
    assert page_number_padding(model, {}, {"kind": "odd", "direction": "right", "chars": 1}) == []
    assert any("第2节首页独立页脚" in note for note in model.unchecked_parts)


def test_compilation_requires_odd_even_setting_and_no_render(tmp_path):
    odd = {"id": "page_number.odd_padding", "scope": {"part": "footer-all"},
           "expected": {"direction": "right", "value": 1, "unit": "char"},
           "enabled": True}
    switch = {"id": "page_number.different_odd_even", "scope": {"part": "footer-all"},
              "expected": {"value": True}, "enabled": True}
    import pytest
    from document_checker.rules_schema import InvalidRuleError
    with pytest.raises(InvalidRuleError, match="奇偶页不同"):
        compile_rules([odd], "custom", "测试")
    _, compiled = compile_rules([switch, odd], "custom", "测试")
    assert [r["checker"] for r in compiled] == ["odd_even_setting", "page_number_padding"]
    model = parse_document(str(_doc(tmp_path)))
    assert page_number_padding(model, {}, compiled[1]["params"]) == []
