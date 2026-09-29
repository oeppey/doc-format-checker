# -*- coding: utf-8 -*-
"""人工评测样本生成脚本（tests/fixtures/manual/）。

设计原则：
- 每份 DOCX 只针对一个规则族制造一处差异，其余部分与 Golden 完全一致，
  这样 gold.json 里「不应出现其他 Finding」的严格断言才有意义。
- 样本由本脚本生成以保证可复现；gold.json 是人工标注，不自动生成。
- 02 的第二个 run 错误已进入正式 Finding 期望。
- 新增反例分别覆盖：密级移到后文、PAGE 域数字为空、附件仅一全角空格。

用法（仓库根目录）：
    .\\.venv\\Scripts\\python tests\\fixtures\\build_manual_fixtures.py
"""
from __future__ import annotations

import copy
import os
import sys

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", "src"))

from document_checker.cli import RULES_DIR  # noqa: E402
from document_checker.synth.golden import build_cfg  # noqa: E402
from document_checker.synth.sample_gen import (  # noqa: E402
    DEFAULT_CONTENT,
    _set_exact_line,
    _set_font,
    build_doc,
)

OUT_ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "manual")


def _out(category: str, filename: str) -> str:
    d = os.path.join(OUT_ROOT, category)
    os.makedirs(d, exist_ok=True)
    return os.path.join(d, filename)


def _base_cfg():
    return build_cfg(RULES_DIR, "dazi")


# ---------------------------------------------------------------- 各样本构建
def build_01_body_font_ok() -> str:
    """正例：正文字体字号完全合规（整份文档零 Finding）。"""
    path = _out("font", "01_body_font_ok.docx")
    build_doc(_base_cfg(), path)
    return path


def build_02_body_font_second_run_bad() -> str:
    """反例：第 4 段正文拆成两个 run，第二个 run 中文字体错为宋体。

    """
    path = _out("font", "02_body_font_second_run_bad.docx")
    build_doc(_base_cfg(), path)
    doc = Document(path)
    p = doc.paragraphs[3]  # 第 4 段（1 起）：首个正文段
    r0 = p.runs[0]
    full = r0.text
    cut = 12
    r0.text = full[:cut]
    r1 = p.add_run(full[cut:])
    r1pr = r1._element.get_or_add_rPr()
    for child in r0._element.rPr:
        r1pr.append(copy.deepcopy(child))
    r1pr.find(qn("w:rFonts")).set(qn("w:eastAsia"), "宋体")  # 仅第二个 run 字体错误
    doc.save(path)
    return path


def build_03_body_font_style_inherited() -> str:
    """正例：第 6 段正文的字体/字号只设在段落样式上，run 不显式设置。

    验证解析器的样式继承兜底（F-01 关注点）；run 上仍显式保留加粗与字符间距。
    """
    path = _out("font", "03_body_font_style_inherited.docx")
    build_doc(_base_cfg(), path)
    doc = Document(path)
    style = doc.styles.add_style("ManualBody", WD_STYLE_TYPE.PARAGRAPH)
    style.font.size = Pt(18)
    srpr = style.element.get_or_add_rPr()
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:eastAsia"), "方正仿宋_GBK")
    srpr.append(rf)
    p = doc.paragraphs[5]  # 第 6 段（1 起）：第二个正文段
    p.style = style
    rpr = p.runs[0]._element.rPr
    for tag in ("w:rFonts", "w:sz", "w:szCs"):
        el = rpr.find(qn(tag))
        if el is not None:
            rpr.remove(el)
    doc.save(path)
    return path


def build_04_page_number_ok() -> str:
    """正例：页脚「— 1 —」居中、宋体四号、自动 PAGE 域。"""
    path = _out("page", "04_page_number_ok.docx")
    build_doc(_base_cfg(), path)
    return path


def build_05_page_number_empty() -> str:
    """反例：整篇没有页脚（页码缺失；并非 PAGE 域数字为空）。"""
    path = _out("page", "05_page_number_empty.docx")
    build_doc(_base_cfg(), path)
    doc = Document(path)
    doc.sections[0].footer.is_linked_to_previous = True  # 删除页脚部件
    doc.save(path)
    return path


def build_18_page_field_empty_result() -> str:
    """反例：保留 PAGE 域，但删除它在 DOCX 中缓存的数字结果。"""
    path = _out("page", "18_page_field_empty_result.docx")
    build_doc(_base_cfg(), path)
    doc = Document(path)
    footer = doc.sections[0].footer.paragraphs[0]
    fields = footer._p.findall(".//" + qn("w:fldSimple"))
    assert len(fields) == 1
    results = fields[0].findall(".//" + qn("w:t"))
    assert len(results) == 1 and results[0].text == "1"
    results[0].text = ""
    doc.save(path)
    return path


def build_06_secrecy_ok() -> str:
    """正例：秘级「秘 密」位于首个非空段、顶格居左、黑体小二加粗、两字间空一格。"""
    path = _out("secrecy", "06_secrecy_ok.docx")
    build_doc(_base_cfg(), path)
    return path


def build_07_secrecy_misplaced() -> str:
    """反例：秘级居中（位置仍在首段；不代表移到后文）。"""
    path = _out("secrecy", "07_secrecy_misplaced.docx")
    build_doc(_base_cfg(), path)
    doc = Document(path)
    doc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(path)
    return path


def build_17_secrecy_after_title() -> str:
    """反例：把密级段移到标题之后，保留密级原有文字和格式。"""
    path = _out("secrecy", "17_secrecy_after_title.docx")
    build_doc(_base_cfg(), path)
    doc = Document(path)
    secrecy, title = doc.paragraphs[:2]
    title._p.addnext(secrecy._p)
    doc.save(path)
    return path


def build_10_heading_number_in_body_bad() -> str:
    """反例：第 11 段正文以「1、」开头（一级标题误用数字+顿号）。"""
    path = _out("heading", "10_heading_number_in_body_bad.docx")
    content = copy.deepcopy(DEFAULT_CONTENT)
    content["blocks"] = content["blocks"] + [("body", "1、专项督导工作推进情况")]
    build_doc(_base_cfg(), path, content=content)
    return path


def build_08_attachment_indent_ok() -> str:
    """正例：附件说明左空两字、正文下空一行、名称后无标点。"""
    path = _out("attachment", "08_attachment_indent_ok.docx")
    build_doc(_base_cfg(), path)
    return path


def build_09_attachment_indent_one_space() -> str:
    """反例：附件说明只留一个半角空格（应左空两字）。"""
    path = _out("attachment", "09_attachment_indent_one_space.docx")
    build_doc(_base_cfg(), path)
    doc = Document(path)
    p = doc.paragraphs[11]  # 第 12 段（1 起）：附件说明
    p.runs[0].text = " 附件：1.XX专项工作实施方案"
    doc.save(path)
    return path


def build_19_attachment_one_ideographic_space() -> str:
    """反例：附件说明只留一个全角空格（一个汉字宽）。"""
    path = _out("attachment", "19_attachment_one_ideographic_space.docx")
    build_doc(_base_cfg(), path)
    doc = Document(path)
    doc.paragraphs[11].runs[0].text = "　附件：1.XX专项工作实施方案"
    doc.save(path)
    return path


def build_11_signature_date_format_bad() -> str:
    """反例：成文日期写作 2026.9.18（应为阿拉伯数字全称 2026年9月18日）。"""
    path = _out("signature", "11_signature_date_format_bad.docx")
    content = copy.deepcopy(DEFAULT_CONTENT)
    content["date"] = "2026.9.18"
    build_doc(_base_cfg(), path, content=content)
    return path


def build_12_signature_blank_lines_bad() -> str:
    """反例：署名与附件说明之间只空一行（应下空二行，warning 级）。"""
    path = _out("signature", "12_signature_blank_lines_bad.docx")
    build_doc(_base_cfg(), path)
    doc = Document(path)
    blank = doc.paragraphs[12]  # 署名前的两个空段之一
    blank._p.getparent().remove(blank._p)
    doc.save(path)
    return path


def build_13_xiaozi_basic_ok() -> str:
    """正例：小字版人工正例（整份文档在小字版规则下零 Finding）。"""
    path = _out("type", "13_xiaozi_basic_ok.docx")
    build_doc(build_cfg(RULES_DIR, "xiaozi"), path)
    return path


def build_14_heading1_misnumbered_labeled() -> str:
    """反例：第 5 段是一级标题（黑体小二加粗，人工标 heading1）但编号错用「1、」。

    角色靠字体启发式归对（黑体→heading1），编号形式由 heading_number 检查器报错。
    """
    path = _out("heading", "14_heading1_misnumbered_labeled.docx")
    content = copy.deepcopy(DEFAULT_CONTENT)
    content["blocks"] = (
        content["blocks"][:2]
        + [("h1", "1、重点工作推进情况")]
        + content["blocks"][2:]
    )
    build_doc(_base_cfg(), path, content=content)
    return path


def build_15_body_in_table() -> str:
    """反例：表格单元格内正文故意使用宋体，检验已解析但未套用正文规则。"""
    path = _out("table", "15_body_in_table.docx")
    build_doc(_base_cfg(), path)
    doc = Document(path)
    table = doc.add_table(rows=1, cols=1)
    cp = table.cell(0, 0).paragraphs[0]
    _set_exact_line(cp, 32.0)
    _set_font(cp.add_run("表格内正文：专项资金拨付使用情况如下。"),
              "宋体", 18, True, 0.4)
    doc.save(path)
    return path


def build_16_multi_section_odd_even_footer() -> str:
    """反例：两节文档，首页页脚与奇偶页页脚不同，且首页/偶数页页码为手打（无 PAGE 域）。

    首页与偶数页页脚已进入 IR 和页码检查；此样本固定其回归行为。
    默认（奇数页）页脚保持合规，其余格式与 Golden 一致。
    """
    from docx.enum.section import WD_SECTION

    path = _out("section", "16_multi_section_odd_even_footer.docx")
    build_doc(_base_cfg(), path)
    doc = Document(path)

    # 先记下署名/日期及其前两个空段，加节加段落后要把它们移回文末
    tail = list(doc.paragraphs[-4:])  # 空、空、署名、日期

    sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
    # 第 2 节正文段（仿宋小二加粗，人工标 body）
    p = doc.add_paragraph()
    _set_exact_line(p, 32.0)
    _set_font(p.add_run("第二节正文：相关工作已按序推进。"),
              "方正仿宋_GBK", 18, True, 0.4)

    body = doc.element.body
    sect_pr = body.find(qn("w:sectPr"))
    for tp in tail:
        sect_pr.addprevious(tp._p)  # lxml 插入即移动，保持日期仍在全文末段

    # 页脚设置：奇偶页不同 + 首页不同；默认（奇数页）页脚沿用合规的「— 1 —」
    doc.settings.odd_and_even_pages_header_footer = True
    sec1 = doc.sections[0]
    sec1.different_first_page_header_footer = True

    fpf = sec1.first_page_footer
    fpf.is_linked_to_previous = False
    _set_font(fpf.paragraphs[0].add_run("1"), "宋体", 14, False, ascii_font="宋体")
    # 首页页脚：手打页码、无 PAGE 域、居左（字体字号合规，缺陷仅页码域与对齐）

    epf = sec1.even_page_footer
    epf.is_linked_to_previous = False
    ep = epf.paragraphs[0]
    ep.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(ep.add_run("— 2 —"), "宋体", 14, False, ascii_font="宋体")
    # 偶数页页脚：横线对、居中，但仍为手打、无 PAGE 域

    doc.save(path)
    return path


BUILDERS = [
    build_01_body_font_ok,
    build_02_body_font_second_run_bad,
    build_03_body_font_style_inherited,
    build_04_page_number_ok,
    build_05_page_number_empty,
    build_18_page_field_empty_result,
    build_06_secrecy_ok,
    build_07_secrecy_misplaced,
    build_17_secrecy_after_title,
    build_10_heading_number_in_body_bad,
    build_08_attachment_indent_ok,
    build_09_attachment_indent_one_space,
    build_19_attachment_one_ideographic_space,
    build_11_signature_date_format_bad,
    build_12_signature_blank_lines_bad,
    build_13_xiaozi_basic_ok,
    build_14_heading1_misnumbered_labeled,
    build_15_body_in_table,
    build_16_multi_section_odd_even_footer,
]

# 自检时各样本适用的文书类型（缺省大字版）
DOC_TYPES = {
    "13_xiaozi_basic_ok.docx": "chengbaogao_xiaozi",
}


def main():
    paths = [fn() for fn in BUILDERS]
    print(f"已生成 {len(paths)} 份人工样本 → {OUT_ROOT}")
    for p in paths:
        print("  " + os.path.relpath(p, OUT_ROOT))

    # 自检：打印每份样本当前的实际 Finding，供核对 gold.json
    from document_checker.rule_engine import RuleEngine
    engine = RuleEngine(RULES_DIR)
    print("\n--- 当前检查器实际输出（用于核对 gold.json）---")
    for p in paths:
        doc_type = DOC_TYPES.get(os.path.basename(p), "chengbaogao_dazi")
        report = engine.check(p, doc_type)
        rel = os.path.relpath(p, OUT_ROOT)
        if not report.findings:
            print(f"{rel}: 无 Finding")
        for f in report.findings:
            print(f"{rel}: [{f.severity}] {f.rule_id} @ {f.location}：{f.message}"
                  f"（期望：{f.expected}；实际：{f.actual}）")


if __name__ == "__main__":
    main()
