"""Independent regression cases for coverage and error reporting."""
from __future__ import annotations

import json

import pytest
from docx import Document

from document_checker.cli import RULES_DIR
from document_checker.report import Report, RuleResult
from document_checker.rule_engine import RuleEngine
from document_checker.rules_schema import InvalidRuleError
from document_checker.typo.corrector import LLMCorrector, _parse_and_validate
from document_checker.typo.detector import SuspectResult, error_label_index
from document_checker.typo.filters import should_skip
from document_checker.typo.pipeline import run_typo_check
from document_checker.typo.splitter import split_paragraph


def _docx(tmp_path, text):
    path = tmp_path / "case.docx"
    doc = Document()
    doc.add_paragraph(text)
    doc.save(path)
    return str(path)


def test_unrecognised_document_type_does_not_fall_back(tmp_path):
    path = _docx(tmp_path, "没有明确版式的测试文件")
    with pytest.raises(ValueError, match="无法判定文书类型"):
        RuleEngine(RULES_DIR).check(path)


@pytest.mark.parametrize("checker,severity", [
    ("does_not_exist", "error"),
    (["font_format"], "error"),
    ("font_format", ["error"]),
])
def test_bad_yaml_rejected_at_load(tmp_path, checker, severity):
    data = {"meta": {"id": "bad"}, "rules": [
        {"id": "broken", "checker": checker, "severity": severity,
         "params": {"role": "body", "size_pt": 16}}
    ]}
    (tmp_path / "bad.yaml").write_text(
        __import__("yaml").safe_dump(data, allow_unicode=True), encoding="utf-8")
    with pytest.raises(InvalidRuleError, match="bad.yaml"):
        RuleEngine(str(tmp_path))


@pytest.mark.parametrize("results,unchecked,expected", [
    ([RuleResult("r", "R", "通过", 0)], [], "通过"),
    ([RuleResult("r", "R", "发现问题", 1)], [], "发现问题"),
    ([RuleResult("r", "R", "通过", 0)], ["表格"], "未检查"),
    ([RuleResult("r", "R", "运行失败", 0)], [], "运行失败"),
])
def test_report_status(results, unchecked, expected):
    assert Report("f", "t", "T", results, unchecked_parts=unchecked).status == expected


def test_sentence_offsets_and_filter_boundary():
    text = "  第一段文字有错误。第二句也有错误！"
    sentences = split_paragraph(0, text)
    assert [(s.start, s.end, s.text) for s in sentences] == [
        (0, 11, "  第一段文字有错误。"),
        (11, len(text), "第二句也有错误！"),
    ]
    assert should_skip("这是一个正常句子。") is None
    assert should_skip("短") == "过短"
    assert should_skip("1234abcd") == "非中文为主"


def test_long_clause_offsets():
    text = "这是一段长句" + "内容" * 34 + "，后半段有错字。"
    sentences = split_paragraph(2, text)
    assert len(sentences) == 2
    assert sentences[0].text == text[:sentences[0].end]
    assert sentences[1].start == sentences[0].end
    assert sentences[1].text == text[sentences[1].start:sentences[1].end]


class _Detector:
    def __init__(self, hints):
        self.hints = hints

    def detect(self, sentences):
        return [SuspectResult(0.9, self.hints) for _ in sentences]


class _Corrector:
    def correct(self, sentence, suspect_chars):
        return [{"原文": "会和", "改为": "会合", "理由": "错字"}]


def test_repeated_fragment_uses_unique_detector_hint(tmp_path):
    path = _docx(tmp_path, "今天会和，明天会和。")
    report = run_typo_check(
        path, detector_backend=(_Detector([(7, "会", 0.9)]), "electra:test"),
        corrector_backend=(_Corrector(), "llm:test"))
    assert len(report.findings) == 1
    assert report.findings[0].sent_start == 7
    assert report.findings[0].location == "第1段·第8字"


def test_repeated_fragment_without_hint_is_failure(tmp_path):
    path = _docx(tmp_path, "今天会和，明天会和。")
    report = run_typo_check(
        path, detector_backend=(_Detector([]), "electra:test"),
        corrector_backend=(_Corrector(), "llm:test"))
    assert report.status == "运行失败"
    assert report.failed_sentences == 1
    assert report.findings == []
    assert any("无法唯一定位" in part for part in report.unchecked_parts)
    assert not any("主题字体" in part for part in report.unchecked_parts)


def test_backend_fallback_never_reports_full_pass(tmp_path):
    path = _docx(tmp_path, "今天我们开展文档检查工作。")
    class CleanDetector:
        def detect(self, sentences):
            return [None for _ in sentences]
    report = run_typo_check(path, detector_backend=(CleanDetector(), "heuristic"),
                            corrector_backend=(_Corrector(), "heuristic"))
    assert report.status == "未检查"
    assert report.to_dict()["status"] == "未检查"
    assert "未检查范围" in report.render_markdown()


def test_corrector_error_counted(tmp_path):
    class FailingCorrector:
        def correct(self, sentence, suspect_chars):
            raise RuntimeError("service down")
    path = _docx(tmp_path, "今天我们开展文档检查工作。")
    report = run_typo_check(
        path, detector_backend=(_Detector([]), "electra:test"),
        corrector_backend=(FailingCorrector(), "llm:test"))
    assert report.status == "运行失败"
    assert report.failed_sentences == 1
    assert "service down" in report.render_markdown()


def test_correction_validator_rejects_rewrite_and_protected_overlap():
    with pytest.raises(ValueError, match="单处错别字"):
        _parse_and_validate(json.dumps([{"原文": "今天开展工作", "改为": "明天结束工作"}]), "今天开展工作")
    with pytest.raises(ValueError, match="受保护"):
        _parse_and_validate(json.dumps([{"原文": "起诉", "改为": "启诉"}]), "决定不起诉")
    assert _parse_and_validate('[{"原文":"会和","改为":"会合"}]', "今天会和。")[0]["改为"] == "会合"


def test_llm_requires_explicit_model(monkeypatch):
    monkeypatch.delenv("DOCUMENT_CHECKER_CORRECTOR_MODEL", raising=False)
    with pytest.raises(ValueError, match="未指定精检模型"):
        LLMCorrector()


def test_electra_label_must_be_explicit():
    assert error_label_index({0: "O", 1: "ERR"}) == 1
    with pytest.raises(ValueError, match="无法唯一确定"):
        error_label_index({0: "LABEL_0", 1: "LABEL_1"})


def test_empty_page_field_result_is_flagged(tmp_path):
    from docx.oxml.ns import qn
    from document_checker.synth.golden import build_cfg
    from document_checker.synth.sample_gen import build_doc
    path = tmp_path / "empty_page.docx"
    build_doc(build_cfg(RULES_DIR, "dazi"), str(path))
    doc = Document(path)
    footer = doc.sections[0].footer.paragraphs[0]
    for text_element in footer._p.findall(".//" + qn("w:t")):
        if text_element.text == "1":
            text_element.text = ""
    doc.save(path)
    report = RuleEngine(RULES_DIR).check(str(path), "chengbaogao_dazi")
    assert any(f.rule_id == "page-number" and f.message == "页码格式不符合要求"
               for f in report.findings)


def test_numpages_field_does_not_count_as_page_number(tmp_path):
    from docx.oxml.ns import qn
    from document_checker.synth.golden import build_cfg
    from document_checker.synth.sample_gen import build_doc
    path = tmp_path / "numpages.docx"
    build_doc(build_cfg(RULES_DIR, "dazi"), str(path))
    doc = Document(path)
    footer = doc.sections[0].footer.paragraphs[0]
    for instruction in footer._p.findall(".//" + qn("w:instrText")):
        instruction.text = (instruction.text or "").replace("PAGE", "NUMPAGES")
    for field in footer._p.findall(".//" + qn("w:fldSimple")):
        field.set(qn("w:instr"), (field.get(qn("w:instr")) or "").replace("PAGE", "NUMPAGES"))
    doc.save(path)
    report = RuleEngine(RULES_DIR).check(str(path), "chengbaogao_dazi")
    assert any(f.rule_id == "page-number" and "未检测到自动页码域" in f.message
               for f in report.findings)


def test_effective_run_style_source_and_implicit_nonbold(tmp_path):
    from docx.enum.style import WD_STYLE_TYPE
    from docx.shared import Pt
    path = tmp_path / "styles.docx"
    doc = Document()
    base = doc.styles.add_style("CaseBase", WD_STYLE_TYPE.PARAGRAPH)
    base.font.size = Pt(18)
    base.font.name = "宋体"
    child = doc.styles.add_style("CaseChild", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = base
    para = doc.add_paragraph(style=child)
    para.add_run("正文")
    doc.save(path)
    from document_checker.docx_parser import parse_document
    run = parse_document(str(path)).paragraphs[0].runs[0]
    assert run.size_pt == 18
    assert run.sources["size_pt"] == "paragraph_style:CaseBase"
    assert run.bold is False
    assert run.sources["bold"] == "implicit_default"


def test_linked_footer_metadata(tmp_path):
    from docx.enum.section import WD_SECTION_START
    from document_checker.docx_parser import parse_document
    path = tmp_path / "linked.docx"
    doc = Document()
    doc.sections[0].footer.paragraphs[0].text = "首页脚"
    section2 = doc.add_section(WD_SECTION_START.NEW_PAGE)
    assert section2.footer.is_linked_to_previous
    doc.save(path)
    model = parse_document(str(path))
    assert model.story_links["section2/default/footer"] == "section1/default/footer"
    assert any(p.section_index == 1 and p.text == "首页脚" for p in model.footer_paras)


@pytest.mark.parametrize("date_text,invalid", [
    ("2026年9月18日", False),
    ("2026.9.18", True),
    ("2026-09-18", True),
    ("2026年09月18日", True),
    ("2026年9月08日", True),
    ("2026年2月30日", True),
])
def test_signature_date_format_from_standard(tmp_path, date_text, invalid):
    from document_checker.synth.golden import build_cfg
    from document_checker.synth.sample_gen import build_doc

    path = tmp_path / "dated.docx"
    build_doc(build_cfg(RULES_DIR, "dazi"), str(path))
    doc = Document(path)
    doc.paragraphs[-1].runs[0].text = date_text
    doc.save(path)

    report = RuleEngine(RULES_DIR).check(str(path), "chengbaogao_dazi")
    date_errors = [f for f in report.findings
                   if f.rule_id == "signature-layout" and "成文日期写法" in f.message]
    assert len(date_errors) == int(invalid)
    assert not [f for f in report.findings if f not in date_errors]


@pytest.mark.parametrize("doc_type,expected_line", [("dazi", 32.0), ("xiaozi", 29.5)])
def test_signature_blank_paragraph_spacing_proxy(tmp_path, doc_type, expected_line):
    from document_checker.docx_parser import parse_document
    from document_checker.synth.golden import build_cfg, resolve_doc_type
    from document_checker.synth.sample_gen import build_doc, _set_exact_line

    path = tmp_path / "blank_spacing.docx"
    build_doc(build_cfg(RULES_DIR, doc_type), str(path))
    model = parse_document(str(path))
    assert [p.line_pt for p in model.paragraphs[-4:-2]] == [expected_line, expected_line]
    doc = Document(path)
    _set_exact_line(doc.paragraphs[-4], 20.0)
    doc.save(path)

    report = RuleEngine(RULES_DIR).check(str(path), resolve_doc_type(doc_type))
    warnings = [f for f in report.findings if f.rule_id == "signature-layout"
                and "空段行距" in f.message]
    assert len(warnings) == 1
    assert warnings[0].severity == "warning"
    assert warnings[0].actual == "exact 20.0 磅"
    assert len(report.findings) == 1
