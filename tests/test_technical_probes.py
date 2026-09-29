# -*- coding: utf-8 -*-
"""技术探针：评审意见第 1、2、3、4、8、12 项的专项验证。

与 test_manual_fixtures.py（人工 gold 严格比对）不同，本文件用程序化构造的
最小样本直接断言解析器／检查器／错字链路的技术行为：

- 项 1 表格顺序与位置：正文—表格—正文顺序、table/r/c/p 路径、不漏段不重复段；
- 项 2 页眉页脚：两节／三节、继承／取消继承、首页／奇偶页，生效故事来自哪一节；
- 项 3 样式有效值：字符样式、段落样式链、直接覆盖、docDefaults，断言 sources；
- 项 4 所有 run：第二／第三 run、中西文混排、空 run；
- 项 8 错字边界：Word 原文 → Finding 的完整位置往返、跳过与失败可追踪；
- 项 12 解析边界：主题字体、表格样式、嵌套表格——断言「正确解析」或显式「未检查」。

四处原先的 xfail 已转为正式断言：过滤原因、主题字体未检查、表格样式未检查、
嵌套表格未检查。主题字体另测检查器不会误报字体。
"""
from __future__ import annotations

import copy
import os

import pytest
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from document_checker.cli import RULES_DIR
from document_checker.docx_parser import parse_document
from document_checker.rule_engine import RuleEngine
from document_checker.synth.golden import build_cfg
from document_checker.synth.sample_gen import _set_font, build_doc
from document_checker.typo.corrector import HeuristicCorrector
from document_checker.typo.detector import HeuristicDetector
from document_checker.typo.pipeline import run_typo_check


@pytest.fixture(scope="module")
def engine():
    return RuleEngine(RULES_DIR)


def _dazi(path, content=None):
    build_doc(build_cfg(RULES_DIR, "dazi"), path, content=content)
    return str(path)


def _copy_rpr(src_run, dst_run):
    dst_rpr = dst_run._element.get_or_add_rPr()
    for child in src_run._element.rPr:
        dst_rpr.append(copy.deepcopy(child))


def _set_style_ea(style, font_name):
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:eastAsia"), font_name)


# ================================================================ 项 1：表格顺序与位置
@pytest.fixture()
def table_doc(tmp_path):
    """正文 — 2x2 表格（r2c1 含两段）— 正文。"""
    path = tmp_path / "t1.docx"
    doc = Document()
    doc.add_paragraph("第一段正文。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).paragraphs[0].add_run("甲")
    table.cell(0, 1).paragraphs[0].add_run("乙")
    c = table.cell(1, 0)
    c.paragraphs[0].add_run("丙1")
    c.add_paragraph("丙2")
    table.cell(1, 1).paragraphs[0].add_run("丁")
    doc.add_paragraph("第二段正文。")
    doc.save(path)
    return parse_document(str(path))


def test_t1_body_order_and_table_paths(table_doc):
    assert table_doc.body_order == ["body/p1", "table1", "body/p2"]
    locations = [p.location for p in table_doc.table_paragraphs]
    assert locations == [
        "table1/r1c1/p1", "table1/r1c2/p1",
        "table1/r2c1/p1", "table1/r2c1/p2", "table1/r2c2/p1",
    ]


def test_t1_no_missing_or_duplicate_paragraphs(table_doc):
    # 表格内容不混入正文段落列表，正文段落不少
    assert [p.text for p in table_doc.paragraphs] == ["第一段正文。", "第二段正文。"]
    texts = [p.text for p in table_doc.table_paragraphs]
    assert texts == ["甲", "乙", "丙1", "丙2", "丁"]
    locs = [p.location for p in table_doc.table_paragraphs]
    assert len(locs) == len(set(locs)), "table/r/c/p 路径重复"


def test_t1_unchecked_parts_declared(table_doc):
    assert any("表格内 5 个非空段落" in u for u in table_doc.unchecked_parts), \
        "下游应能明确知道哪些表格文字未检查"


# ================================================================ 项 2：页眉页脚
def test_t2_two_section_inherited_footer(tmp_path, engine):
    """两节；第 2 节页脚继承第 1 节 → 不重复解析、不重复报错，链接可追溯。"""
    path = _dazi(tmp_path / "t2a.docx")
    doc = Document(path)
    doc.add_section(WD_SECTION.NEW_PAGE)
    header = doc.sections[0].header
    header.is_linked_to_previous = False
    header.paragraphs[0].add_run("内部资料")
    doc.save(path)

    model = parse_document(path)
    assert model.story_links.get("section2/default/footer") == "section1/default/footer"
    assert len(model.footer_paras) == 1, "继承页脚不应重复解析"
    assert [p.text for p in model.header_paras] == ["内部资料"]
    assert any("页眉" in u for u in model.unchecked_parts), "页眉已解析但规则未应用应显式声明"

    report = engine.check(path, "chengbaogao_dazi")
    assert report.findings == [], f"继承场景不应产生 Finding：{[f.message for f in report.findings]}"


def test_t2_three_section_middle_override(tmp_path, engine):
    """三节；第 2 节取消继承并改成手打页码，第 3 节继承 → 只报第 2 节，不漏不重。"""
    path = _dazi(tmp_path / "t2b.docx")
    doc = Document(path)
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_section(WD_SECTION.NEW_PAGE)
    sec2 = doc.sections[1]
    sec2.footer.is_linked_to_previous = False
    fp = sec2.footer.paragraphs[0]
    _set_font(fp.add_run("2"), "宋体", 14, False)  # 手打页码、居左
    doc.save(path)

    model = parse_document(path)
    assert model.story_links.get("section3/default/footer") == "section2/default/footer"
    assert "section2/default/footer" not in model.story_links, "取消继承的节不应记为链接"

    report = engine.check(path, "chengbaogao_dazi")
    locs = [f.location for f in report.findings]
    assert len(report.findings) == 2, f"应只报第 2 节的 2 条问题，实际：{locs}"
    assert all("第2节" in loc for loc in locs)
    assert not any("第3节" in loc for loc in locs), "继承第 2 节的第 3 节不应重复报错"
    messages = [f.message for f in report.findings]
    assert any("自动页码域" in m for m in messages)
    assert any("未左右居中" in m for m in messages)


def test_t2_first_even_default_kinds(tmp_path, engine):
    """首页／偶数页／普通三类页脚分别解析；页码规则作用于每一类生效页脚。"""
    path = _dazi(tmp_path / "t2c.docx")
    doc = Document(path)
    doc.settings.odd_and_even_pages_header_footer = True
    sec1 = doc.sections[0]
    sec1.different_first_page_header_footer = True
    fpf = sec1.first_page_footer
    fpf.is_linked_to_previous = False
    _set_font(fpf.paragraphs[0].add_run("1"), "宋体", 14, False)  # 手打、居左
    epf = sec1.even_page_footer
    epf.is_linked_to_previous = False
    ep = epf.paragraphs[0]
    ep.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(ep.add_run("— 2 —"), "宋体", 14, False)  # 手打、居中
    doc.save(path)

    model = parse_document(path)
    kinds = sorted(p.kind for p in model.footer_paras)
    assert kinds == ["default", "even", "first"]
    locs = {p.kind: p.location for p in model.footer_paras}
    assert "首页页脚" in locs["first"] and "偶数页页脚" in locs["even"] and "普通页脚" in locs["default"]

    report = engine.check(path, "chengbaogao_dazi")
    first_hits = [f for f in report.findings if "首页页脚" in f.location]
    even_hits = [f for f in report.findings if "偶数页页脚" in f.location]
    default_hits = [f for f in report.findings if "普通页脚" in f.location]
    assert len(first_hits) == 2, "首页页脚：无 PAGE 域 + 未居中"
    assert len(even_hits) == 1, "偶数页页脚：无 PAGE 域"
    assert default_hits == [], "合规的普通页脚不应报错"


def test_t2_kinds_absent_when_disabled(tmp_path):
    """未开启首页不同／奇偶页不同时，不应解析出对应故事。"""
    model = parse_document(_dazi(tmp_path / "t2d.docx"))
    assert {p.kind for p in model.footer_paras} == {"default"}


# ================================================================ 项 3：样式有效值
def test_t3_character_style_source(tmp_path):
    """字符样式提供字体/字号；sources 记为 character_style。"""
    path = tmp_path / "t3a.docx"
    doc = Document()
    cs = doc.styles.add_style("MyChar", WD_STYLE_TYPE.CHARACTER)
    _set_style_ea(cs, "黑体")
    cs.font.size = Pt(16)
    p = doc.add_paragraph()
    r = p.add_run("字符样式测试")
    r.style = cs
    doc.save(path)

    run = parse_document(str(path)).paragraphs[0].runs[0]
    assert run.font_ea == "黑体" and run.sources["font_ea"] == "character_style:MyChar"
    assert run.size_pt == 16 and run.sources["size_pt"] == "character_style:MyChar"


def test_t3_paragraph_style_chain_and_direct_override(tmp_path):
    """段落样式链：子样式给字号、父样式给字体；run 直接设置加粗优先。"""
    path = tmp_path / "t3b.docx"
    doc = Document()
    base = doc.styles.add_style("BaseA", WD_STYLE_TYPE.PARAGRAPH)
    _set_style_ea(base, "楷体")
    child = doc.styles.add_style("ChildB", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = base
    child.font.size = Pt(18)
    p = doc.add_paragraph()
    p.style = child
    r = p.add_run("样式链测试")
    r.font.bold = True
    doc.save(path)

    run = parse_document(str(path)).paragraphs[0].runs[0]
    assert run.font_ea == "楷体" and run.sources["font_ea"] == "paragraph_style:BaseA"
    assert run.size_pt == 18 and run.sources["size_pt"] == "paragraph_style:ChildB"
    assert run.bold is True and run.sources["bold"] == "run"


def test_t3_direct_run_beats_style(tmp_path):
    """反例对照：run 直接格式覆盖段落样式。"""
    path = tmp_path / "t3c.docx"
    doc = Document()
    ps = doc.styles.add_style("StyleC", WD_STYLE_TYPE.PARAGRAPH)
    _set_style_ea(ps, "宋体")
    p = doc.add_paragraph()
    p.style = ps
    r = p.add_run("直接覆盖测试")
    r._element.get_or_add_rPr()
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:eastAsia"), "仿宋")
    r._element.rPr.append(rf)
    doc.save(path)

    run = parse_document(str(path)).paragraphs[0].runs[0]
    assert run.font_ea == "仿宋" and run.sources["font_ea"] == "run"


def test_t3_doc_defaults_source(tmp_path):
    """run／样式都没设时落到 docDefaults，sources 记为 doc_defaults。"""
    path = tmp_path / "t3d.docx"
    doc = Document()
    rpr = doc.styles.element.find(qn("w:docDefaults") + "/" + qn("w:rPrDefault") + "/" + qn("w:rPr"))
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    rf.set(qn("w:eastAsia"), "仿宋_GBK")
    doc.add_paragraph().add_run("默认层测试")
    doc.save(path)

    run = parse_document(str(path)).paragraphs[0].runs[0]
    assert run.font_ea == "仿宋_GBK" and run.sources["font_ea"] == "doc_defaults"


# ================================================================ 项 4：所有 run
def test_t4_second_and_third_run_checked(tmp_path, engine):
    """一段三 run：run2 字体错、run3 字号错 → 两条 Finding 且定位到 run。"""
    path = _dazi(tmp_path / "t4a.docx")
    doc = Document(path)
    p = doc.paragraphs[3]
    r0 = p.runs[0]
    full = r0.text
    t1, t2, t3 = full[:8], full[8:16], full[16:]
    r0.text = t1
    r1 = p.add_run(t2)
    _copy_rpr(r0, r1)
    r1._element.rPr.find(qn("w:rFonts")).set(qn("w:eastAsia"), "宋体")
    r2 = p.add_run(t3)
    _copy_rpr(r0, r2)
    r2.font.size = Pt(16)
    doc.save(path)

    report = engine.check(path, "chengbaogao_dazi")
    msgs = [(f.location, f.message) for f in report.findings]
    assert len(report.findings) == 2, f"应恰好 2 条 Finding：{msgs}"
    assert any("run 2" in loc and "中文字体" in m for loc, m in msgs)
    assert any("run 3" in loc and "字号" in m for loc, m in msgs)


def test_t4_mixed_cjk_ascii_run(tmp_path, engine):
    """中西文混排 run：中文合格、西文字体错 → 报西文字体。"""
    path = _dazi(tmp_path / "t4b.docx")
    doc = Document(path)
    p = doc.paragraphs[3]
    r = p.runs[0]
    r.text = "2026年计划完成率105%"
    r.font.name = "Times New Roman"
    doc.save(path)

    report = engine.check(path, "chengbaogao_dazi")
    msgs = [(f.location, f.message) for f in report.findings]
    assert len(report.findings) == 1, f"应恰好 1 条 Finding：{msgs}"
    assert "西文字体" in msgs[0][1] and "run 1" in msgs[0][0]


def test_t4_empty_and_space_runs_skipped(tmp_path, engine):
    """空 run 与纯空格 run：不报错、不崩溃。"""
    path = _dazi(tmp_path / "t4c.docx")
    doc = Document(path)
    p = doc.paragraphs[3]
    r0 = p.runs[0]
    full = r0.text
    r0.text = full[:8]
    r_empty = p.add_run("")
    _copy_rpr(r0, r_empty)
    r_space = p.add_run(" ")
    _copy_rpr(r0, r_space)
    r_rest = p.add_run(full[8:])
    _copy_rpr(r0, r_rest)
    doc.save(path)

    report = engine.check(path, "chengbaogao_dazi")
    assert report.findings == [], f"空 run／空格 run 不应产生 Finding：{[f.message for f in report.findings]}"


# ================================================================ 项 8：错字边界
def _heuristic_report(path):
    return run_typo_check(
        str(path),
        detector_backend=(HeuristicDetector(), "heuristic"),
        corrector_backend=(HeuristicCorrector(), "heuristic"),
    )


def test_t8_word_to_finding_location_roundtrip(tmp_path):
    """Word 原文 → 分句 → Finding：段落号、句内偏移、所在句全部一致。"""
    path = _dazi(tmp_path / "t8a.docx")
    doc = Document(path)
    p = doc.paragraphs[3]
    p.runs[0].text = p.runs[0].text.replace("部署", "布署", 1)
    doc.save(path)

    report = _heuristic_report(path)
    assert len(report.findings) == 1
    f = report.findings[0]
    para_text = parse_document(path).paragraphs[f.para_index - 1].text
    assert para_text[f.sent_start: f.sent_start + len(f.original)] == f.original, \
        "Finding 偏移指回原文必须是错字本身"
    assert (f.original, f.suggestion) == ("布署", "部署")
    assert f.original in f.sentence and f.sentence in para_text
    assert f.sentence.endswith("。"), "短段不应触发按逗号的二次切分"


def test_t8_long_paragraph_clause_offset(tmp_path):
    """长句按逗号二次切分后，跨子句的偏移仍然指回原文正确位置。"""
    from document_checker.synth.sample_gen import DEFAULT_CONTENT
    content = copy.deepcopy(DEFAULT_CONTENT)
    long_para = ("今年以来，我室按照委机关统一布署和年度工作要点，扎实推进各项专项工作，"
                 "强化跟踪问效和督导检查，确保高质量完成年度目标任务，各项工作取得阶段性成效。")
    content["blocks"] = [("h1", "一、工作开展情况"), ("body", long_para)]
    path = _dazi(tmp_path / "t8b.docx", content=content)

    report = _heuristic_report(path)
    assert len(report.findings) == 1
    f = report.findings[0]
    para_text = parse_document(path).paragraphs[f.para_index - 1].text
    assert len(para_text) > 60, "样本必须是触发二次切分的长段"
    assert para_text[f.sent_start: f.sent_start + len(f.original)] == f.original
    assert f.sentence.endswith("，"), "错字在段落中部，所在句应是逗号切出的子句"


def test_t8_skip_reasons_and_status(tmp_path):
    """过滤逐句计数、跳过总量与状态；专用跳过原因应可追踪。"""
    path = tmp_path / "t8c.docx"
    doc = Document()
    for text in ("泰纪发〔2026〕5号", "2026年9月18日", "— 3 —",
                 "11010119900307775X", "This is an English sentence.", "好的",
                 "这是需要检查的正常句子，内容长度足够通过过滤。"):
        doc.add_paragraph(text)
    doc.save(path)

    report = _heuristic_report(path)
    assert report.skip_reasons == {
        "文号行": 1, "日期行": 1, "页码行": 1, "证件号码": 1,
        "非中文为主": 1, "过短": 1,
    }
    assert report.skipped == 6 and report.total_sentences == 7
    assert report.status == "未检查"
    assert any("ELECTRA 初筛未执行" in u for u in report.unchecked_parts)
    assert any("精检未执行" in u for u in report.unchecked_parts)


def test_t8_specific_skip_reasons_take_precedence(tmp_path):
    """专用模式应优先于 CJK 占比给出准确原因（过滤顺序缺陷）。"""
    path = tmp_path / "t8d.docx"
    doc = Document()
    for text in ("泰纪发〔2026〕5号", "2026年9月18日", "— 3 —", "11010119900307775X"):
        doc.add_paragraph(text)
    doc.save(path)
    report = _heuristic_report(path)
    assert report.skip_reasons == {
        "文号行": 1, "日期行": 1, "页码行": 1, "证件号码": 1,
    }


def test_t8_corrector_failure_tracked(tmp_path):
    """精检异常：按句计数、带位置写进未检查清单，状态为「运行失败」。"""
    class BoomCorrector:
        def correct(self, sentence, suspect_chars=None):
            raise RuntimeError("模拟服务不可用")

    path = _dazi(tmp_path / "t8e.docx")
    doc = Document(path)
    p = doc.paragraphs[3]
    p.runs[0].text = p.runs[0].text.replace("部署", "布署", 1)
    doc.save(path)

    report = run_typo_check(
        path,
        detector_backend=(HeuristicDetector(), "heuristic"),
        corrector_backend=(BoomCorrector(), "llm:boom"),
    )
    # 启发式检测器会把含叠字（XX）和错字的 4 句都判为嫌疑，全部精检失败
    assert report.suspect == 4
    assert report.failed_sentences == report.suspect
    assert report.status == "运行失败"
    assert any("精检失败" in u and "第4段" in u for u in report.unchecked_parts), \
        "失败应带段落位置可追踪"


# ================================================================ 项 12：解析边界
def test_t12_theme_font_not_guessed(tmp_path):
    """主题字体（eastAsiaTheme）：解析器不猜，值与来源都为空。"""
    path = tmp_path / "t12a.docx"
    doc = Document()
    r = doc.add_paragraph().add_run("主题字体测试")
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:eastAsiaTheme"), "majorEastAsia")
    r._element.get_or_add_rPr().append(rf)
    doc.save(path)

    run = parse_document(str(path)).paragraphs[0].runs[0]
    assert run.font_ea is None and "font_ea" not in run.sources


def test_t12_theme_font_declared_unchecked(tmp_path):
    """主题字体应显式列入未检查清单。"""
    path = tmp_path / "t12a2.docx"
    doc = Document()
    r = doc.add_paragraph().add_run("主题字体测试")
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:eastAsiaTheme"), "majorEastAsia")
    r._element.get_or_add_rPr().append(rf)
    doc.save(path)
    model = parse_document(str(path))
    assert any("主题字体" in un for un in model.unchecked_parts)


def test_t12_table_style_not_applied_silently(tmp_path):
    """表格样式携带的字体当前不应用于单元格 run（记录现状，不靠猜测）。"""
    path = tmp_path / "t12b.docx"
    doc = Document()
    ts = doc.styles.add_style("MyTableStyle", WD_STYLE_TYPE.TABLE)
    _set_style_ea(ts, "黑体")
    table = doc.add_table(rows=1, cols=1)
    table.style = ts
    table.cell(0, 0).paragraphs[0].add_run("表格样式测试")
    doc.save(path)

    run = parse_document(str(path)).table_paragraphs[0].runs[0]
    assert run.font_ea is None, "表格样式字体不应静默生效（现状：未实现）"


def test_t12_table_style_declared_unchecked(tmp_path):
    """带样式表格应显式声明「表格样式未解析」。"""
    path = tmp_path / "t12b2.docx"
    doc = Document()
    ts = doc.styles.add_style("MyTableStyle2", WD_STYLE_TYPE.TABLE)
    _set_style_ea(ts, "黑体")
    table = doc.add_table(rows=1, cols=1)
    table.style = ts
    table.cell(0, 0).paragraphs[0].add_run("表格样式测试")
    doc.save(path)
    model = parse_document(str(path))
    assert any("表格样式" in un for un in model.unchecked_parts)


def test_t12_nested_table_currently_dropped(tmp_path):
    """嵌套表格：外层单元格文字解析正常，内层表格文字当前不进 IR（记录现状）。"""
    path = tmp_path / "t12c.docx"
    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    cell = outer.cell(0, 0)
    cell.paragraphs[0].add_run("外层文字")
    nested = cell.add_table(rows=1, cols=1)
    nested.cell(0, 0).paragraphs[0].add_run("嵌套文字")
    doc.save(path)

    model = parse_document(str(path))
    all_text = [p.text for p in model.paragraphs + model.table_paragraphs]
    assert any("外层文字" in t for t in all_text)
    assert not any("嵌套文字" in t for t in all_text), "现状：嵌套表格文字不进 IR"


def test_t12_nested_table_declared_unchecked(tmp_path):
    """嵌套表格应显式声明未检查。"""
    path = tmp_path / "t12c2.docx"
    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    cell = outer.cell(0, 0)
    cell.paragraphs[0].add_run("外层文字")
    nested = cell.add_table(rows=1, cols=1)
    nested.cell(0, 0).paragraphs[0].add_run("嵌套文字")
    doc.save(path)
    model = parse_document(str(path))
    assert any("嵌套表格" in un for un in model.unchecked_parts)


def test_t12_theme_font_is_unchecked_without_false_font_error(tmp_path, engine):
    """Theme font remains unresolved; the same run's other known properties are still checked."""
    path = tmp_path / "theme_in_body.docx"
    _dazi(path)
    doc = Document(path)
    run = doc.paragraphs[3].runs[0]
    fonts = run._element.get_or_add_rPr().find(qn("w:rFonts"))
    fonts.attrib.pop(qn("w:eastAsia"), None)
    fonts.set(qn("w:eastAsiaTheme"), "majorEastAsia")
    doc.save(path)

    model = parse_document(str(path))
    body_run = model.paragraphs[3].runs[0]
    assert body_run.font_ea is None
    assert "font_ea" in body_run.unresolved
    assert any("body/p4/run1 主题字体 font_ea" in part for part in model.unchecked_parts)
    report = engine.check(str(path), "chengbaogao_dazi")
    assert report.status == "未检查"
    assert not [f for f in report.findings if f.rule_id == "font-body"]
    assert not report.findings


def test_t12_theme_inherited_below_explicit_font_is_not_unresolved(tmp_path):
    """A direct run font wins over a theme reference in the paragraph style."""
    path = tmp_path / "explicit_over_theme.docx"
    doc = Document()
    style = doc.styles.add_style("ThemeBase", WD_STYLE_TYPE.PARAGRAPH)
    rpr = style.element.get_or_add_rPr()
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:eastAsiaTheme"), "majorEastAsia")
    rpr.append(rf)
    paragraph = doc.add_paragraph(style=style)
    run = paragraph.add_run("明确字体")
    _set_font(run, "方正仿宋_GBK", 18, True, 0.4)
    doc.save(path)

    parsed = parse_document(str(path)).paragraphs[0].runs[0]
    assert parsed.font_ea == "方正仿宋_GBK"
    assert "font_ea" not in parsed.unresolved
