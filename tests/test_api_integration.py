from pathlib import Path

from fastapi.testclient import TestClient

from document_checker.api import create_app
from document_checker.standardization import CATALOG

ROOT = Path(__file__).resolve().parent
GOOD = ROOT / "fixtures/manual/font/01_body_font_ok.docx"
MIXED = ROOT / "fixtures/manual/font/02_body_font_second_run_bad.docx"


def _upload(client, path, endpoint, data=None):
    with path.open("rb") as fp:
        return client.post(endpoint, data=data or {},
                           files={"file": (path.name, fp,
                                   "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})


def test_catalog_and_real_observations(tmp_path):
    client = TestClient(create_app(data_dir=tmp_path / "templates"))
    response = client.get("/api/catalog")
    assert response.status_code == 200
    items = response.json()["items"]
    assert len(items) == len(CATALOG) == 36
    assert {x["id"]: x["capability"] for x in items}["structural.seal_layout"] == "unavailable"

    good = _upload(client, GOOD, "/api/templates/extract")
    assert good.status_code == 200
    data = good.json()
    assert data["candidates"]
    assert all(not x["enabled"] for x in data["candidates"])
    assert any(x["id"] == "font.cjk" and x["scope"].get("role") == "body"
               for x in data["candidates"])

    mixed = _upload(client, MIXED, "/api/templates/extract")
    assert mixed.status_code == 200
    assert any(x["id"] == "font.cjk" and x["role"] == "body"
               for x in mixed.json()["mixed"])


def test_template_edit_persist_and_review(tmp_path):
    data_dir = tmp_path / "templates"
    client = TestClient(create_app(data_dir=data_dir))
    extracted = _upload(client, GOOD, "/api/templates/extract").json()
    rule = next(x for x in extracted["candidates"]
                if x["id"] == "font.cjk" and x["scope"].get("role") == "body")
    rule["enabled"] = True
    rule["expected"] = {"value": "黑体", "aliases": []}
    payload = {"name": "字体测试模板", "source": GOOD.name, "rules": [rule]}
    validated = client.post("/api/templates/validate", json=payload)
    assert validated.status_code == 200
    assert validated.json()["active_rule_count"] == 1

    saved = client.post("/api/templates", json=payload)
    assert saved.status_code == 200
    template_id = saved.json()["id"]
    assert client.get(f"/api/templates/{template_id}").json()["rules"][0]["expected"]["value"] == "黑体"
    assert any(x["id"] == template_id for x in
               client.get("/api/templates").json()["templates"])

    reviewed = _upload(client, GOOD, "/api/reviews", {"template_id": template_id})
    assert reviewed.status_code == 200
    report = reviewed.json()
    assert any(f["rule_id"] == "font-body" for f in report["format"]["findings"])
    assert report["typo"]["status"] in {"未检查", "发现问题", "运行失败", "通过"}
    assert "preview_status" in report
    assert "repair_status" in report


def test_units_and_unavailable_rules_are_validated(tmp_path):
    client = TestClient(create_app(data_dir=tmp_path / "templates"))
    margin = {
        "id": "page.margin_top", "scope": {},
        "expected": {"value": 38, "unit": "mm"}, "enabled": True,
    }
    response = client.post("/api/templates/validate",
                           json={"name": "单位", "rules": [margin]})
    assert response.status_code == 200
    assert response.json()["normalized_rules"][0]["expected"] == {"value": 3.8, "unit": "cm"}

    seal = {"id": "structural.seal_layout", "scope": {},
            "expected": {"value": True}, "enabled": True}
    response = client.post("/api/templates/validate",
                           json={"name": "印章", "rules": [seal]})
    assert response.status_code == 422
    assert _upload(client, GOOD, "/api/templates/extract").status_code == 200
    response = client.post("/api/templates/extract",
                           files={"file": ("old.doc", b"not-docx")})
    assert response.status_code == 400



def test_multiple_line_spacing_stays_a_ratio(tmp_path):
    from docx import Document
    from docx.enum.text import WD_LINE_SPACING
    from document_checker.docx_parser import parse_document

    doc = Document()
    style = doc.styles["Normal"]
    style.paragraph_format.line_spacing = 32
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    para = doc.add_paragraph("正文倍数行距")
    para.paragraph_format.line_spacing = 1.5
    path = tmp_path / "multiple.docx"
    doc.save(path)
    parsed = parse_document(str(path)).paragraphs[0]
    assert parsed.line_rule == "auto"
    assert parsed.line_multiple == 1.5
    assert parsed.line_pt is None

def test_odd_even_setting_and_padding_options(tmp_path):
    client = TestClient(create_app(data_dir=tmp_path / "templates"))
    catalog = {item["id"]: item for item in client.get("/api/catalog").json()["items"]}
    assert catalog["page_number.odd_padding"]["kind"] == "padding"
    assert catalog["page_number.even_padding"]["capability"] == "needs_render"
    sample = _upload(client, GOOD, "/api/templates/extract").json()
    options = {r["id"]: r for r in sample["candidates"]}
    assert options["page_number.odd_padding"]["expected"] == {
        "direction": "right", "value": 1, "unit": "char"}
    assert options["page_number.even_padding"]["expected"] == {
        "direction": "left", "value": 1, "unit": "char"}
    assert not options["page_number.different_odd_even"]["expected"]["value"]
    switch = options["page_number.different_odd_even"]
    switch["enabled"] = True
    switch["expected"] = {"value": True}
    result = client.post("/api/templates/validate",
                         json={"name": "奇偶页", "rules": [switch]})
    assert result.status_code == 200
    assert result.json()["compiled_rules"][0]["checker"] == "odd_even_setting"
    saved = client.post("/api/templates", json={"name": "奇偶页", "rules": [switch, options["page_number.odd_padding"], options["page_number.even_padding"]]})
    assert len(saved.json()["rules"]) == 3
    assert len(saved.json()["compiled_rules"]) == 1
    assert saved.status_code == 200
    reviewed = _upload(client, GOOD, "/api/reviews",
                       {"template_id": saved.json()["id"]})
    assert reviewed.status_code == 200
    assert any(f["rule_id"] == "odd-even-setting"
               for f in reviewed.json()["format"]["findings"])
    padding = options["page_number.odd_padding"]
    padding["expected"] = {"direction": "none", "value": 1, "unit": "char"}
    invalid = client.post("/api/templates/validate",
                          json={"name": "无效留空", "rules": [padding]})
    assert invalid.status_code == 422
    padding["expected"] = {"direction": "none", "value": 0, "unit": "char"}
    padding["enabled"] = True
    unsupported = client.post("/api/templates/validate",
                              json={"name": "视觉规则", "rules": [padding]})
    assert unsupported.status_code == 422


