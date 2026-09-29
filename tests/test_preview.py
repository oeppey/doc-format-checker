from pathlib import Path

import pymupdf as fitz
from docx import Document
from fastapi.testclient import TestClient

from document_checker.api import create_app
from document_checker.docx_parser import parse_document
from document_checker.preview import PreviewError, map_findings


def test_pdf_character_mapping_for_run_typo_and_unlocated(tmp_path):
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("alpha ")
    paragraph.add_run("bravo")
    paragraph.add_run(" charlie")
    docx = tmp_path / "source.docx"
    doc.save(docx)
    model = parse_document(str(docx))

    pdf_path = tmp_path / "preview.pdf"
    pdf = fitz.open()
    page = pdf.new_page(width=595, height=842)
    page.insert_text((72, 100), "alpha bravo charlie", fontsize=12)
    pdf.save(pdf_path)
    pdf.close()

    found = map_findings(
        pdf_path, model,
        [{"location": "第1段「alpha bravo…」（run 2）"},
         {"location": "第9段「不存在」"},
         {"location": "第1节普通页脚第1段"}],
        [{"para_index": 1, "sent_start": 6, "original": "bravo"}],
        [{"number": 1, "width_pt": 595, "height_pt": 842}],
    )
    assert found[0]["precision"] == "run"
    assert found[0]["rects"][0]["page"] == 1
    assert found[1]["precision"] == "unlocated" and not found[1]["rects"]
    assert found[2]["precision"] == "unlocated" and not found[2]["rects"]
    assert found[3]["precision"] == "text"
    assert found[3]["rects"][0]["x"] == found[0]["rects"][0]["x"]


def test_render_failure_preserves_review_result(tmp_path, monkeypatch):
    from document_checker import api

    def unavailable(*args):
        raise PreviewError("渲染器不存在")

    monkeypatch.setattr(api, "render_docx", unavailable)
    client = TestClient(create_app(data_dir=tmp_path / "templates"))
    source = Path(__file__).parent / "fixtures/manual/font/01_body_font_ok.docx"
    with source.open("rb") as fp:
        response = client.post(
            "/api/reviews", data={"template_id": "builtin:chengbaogao_dazi"},
            files={"file": (source.name, fp,
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
    assert response.status_code == 200
    result = response.json()
    assert result["format"]["findings"] is not None
    assert result["preview"]["status"] == "failed"
    assert result["preview"]["reason"] == "渲染器不存在"
    assert result["preview"]["pages"] == []
    assert client.get("/api/reviews/not-a-uuid/preview.pdf").status_code == 404




def test_real_docx_red_typo_mapping(tmp_path):
    """The rendered PDF, page image and typo coordinates share one geometry."""
    import pytest
    from document_checker.preview import _soffice, render_docx
    from document_checker.typo.pipeline import run_typo_check
    from document_checker.typo.detector import HeuristicDetector
    from document_checker.typo.corrector import HeuristicCorrector

    try:
        _soffice()
    except PreviewError:
        pytest.skip("LibreOffice is not installed")
    doc = Document()
    doc.add_paragraph("今年以来我单位认真布署专项工作，确保各项任务按期完成。")
    source = tmp_path / "typo.docx"
    doc.save(source)
    model = parse_document(str(source))
    report = run_typo_check(
        str(source), model=model,
        detector_backend=(HeuristicDetector(), "heuristic"),
        corrector_backend=(HeuristicCorrector(), "heuristic"),
    )
    assert [f.original for f in report.findings] == ["布署"]
    output = tmp_path / "pages"
    pages = render_docx(source, output)
    mapped = map_findings(output / "preview.pdf", model, [],
                          report.to_dict()["findings"], pages)
    assert len(pages) == 1
    assert mapped[0]["precision"] == "text"
    assert mapped[0]["rects"][0]["page"] == 1
    assert (output / pages[0]["image_name"]).is_file()
